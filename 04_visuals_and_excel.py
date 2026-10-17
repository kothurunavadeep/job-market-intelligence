"""
Step 4 - EDA CHARTS + DASHBOARD PREVIEW + EXCEL INSIGHTS WORKBOOK
Reads SQL results from outputs/sql_results and clean data from data/.
"""
import os
import glob
import json
import pandas as pd
import numpy as np
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.ticker import FuncFormatter
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, Reference

os.makedirs("outputs/charts", exist_ok=True)
R = {os.path.basename(p).split("_")[0]: pd.read_csv(p) for p in sorted(glob.glob("outputs/sql_results/Q*.csv"))}
jobs = pd.read_csv("data/clean_jobs.csv")
log = json.load(open("data/cleaning_log.json"))

# ------------------------------------------------------------------ style
NAVY, TEAL, ORANGE, GREY, LIGHT = "#1F3A5F", "#2A9D8F", "#E76F51", "#6B7280", "#F3F4F6"
plt.rcParams.update({
    "font.family": "DejaVu Sans", "axes.spines.top": False, "axes.spines.right": False,
    "axes.titleweight": "bold", "axes.titlesize": 12, "axes.labelsize": 9, "xtick.labelsize": 8.5,
    "ytick.labelsize": 8.5, "figure.facecolor": "white",
})
fmt_k = FuncFormatter(lambda x, _: f"{x/1000:.0f}K" if x >= 1000 else f"{x:.0f}")


def hbar(ax, labels, values, color=TEAL, title="", xlabel="", fmt="{:,.0f}", highlight=None):
    y = np.arange(len(labels))[::-1]
    cols = [ORANGE if (highlight and l in highlight) else color for l in labels]
    ax.barh(y, values, color=cols, height=0.68)
    ax.set_yticks(y)
    ax.set_yticklabels(labels)
    ax.set_title(title, loc="left", color=NAVY)
    ax.set_xlabel(xlabel, color=GREY)
    mx = max(values)
    for yi, v in zip(y, values):
        ax.text(v + mx * 0.01, yi, fmt.format(v), va="center", fontsize=8, color="#374151")
    ax.set_xlim(0, mx * 1.15)


# ------------------------------------------------------------------ individual charts
data_jobs = jobs[jobs["is_data_role"]]
n_data = len(data_jobs)

# C1 skills in data roles
q5 = R["Q05"].head(15)
fig, ax = plt.subplots(figsize=(8, 5.5))
hbar(ax, q5["skill"], q5["pct_of_data_jobs"], title="Most in-demand skills in Data roles",
     xlabel="% of data job postings listing the skill", fmt="{:.1f}%",
     highlight=["Python", "SQL", "Power BI", "Excel", "Tableau"])
fig.text(0.01, 0.01, f"n = {n_data:,} data postings (Data Analyst / Scientist / Engineer)", fontsize=8, color=GREY)
fig.tight_layout(rect=(0, 0.03, 1, 1)); fig.savefig("outputs/charts/01_top_skills_data_roles.png", dpi=160); plt.close(fig)

# C2 role demand
q2 = R["Q02"]
q2 = q2[~q2["role_family"].isin(["Other", "Admin & General Mgmt"])].head(15)
fig, ax = plt.subplots(figsize=(8, 5.5))
hbar(ax, q2["role_family"], q2["job_postings"], color=NAVY, title="Role demand: job postings by role family",
     xlabel="Job postings", highlight=["Data Analyst", "Data Scientist", "Data Engineer"])
fig.tight_layout(); fig.savefig("outputs/charts/02_role_demand.png", dpi=160); plt.close(fig)

# C3 city demand
q9 = R["Q09"]
q9 = q9[q9["city"] != "Remote / Unspecified"].head(12)
fig, ax = plt.subplots(figsize=(8, 5))
x = np.arange(len(q9))
ax.bar(x, q9["job_postings"], color=NAVY, label="All jobs")
ax.bar(x, q9["tech_jobs"], color=TEAL, label="Tech jobs")
ax.bar(x, q9["data_jobs"], color=ORANGE, label="Data jobs")
ax.set_xticks(x); ax.set_xticklabels(q9["city"], rotation=40, ha="right")
ax.yaxis.set_major_formatter(fmt_k)
ax.set_title("Hiring hubs: postings by city", loc="left", color=NAVY)
ax.legend(frameon=False, fontsize=8)
fig.tight_layout(); fig.savefig("outputs/charts/03_city_demand.png", dpi=160); plt.close(fig)

# C4 salary by role
q13 = R["Q13"]
q13 = q13[q13["jobs_with_salary"] >= 100].sort_values("median_lpa", ascending=False).head(14)
fig, ax = plt.subplots(figsize=(8, 5.5))
hbar(ax, q13["role_family"], q13["median_lpa"], color=TEAL, title="Median salary by role family (LPA)",
     xlabel="Median annual salary midpoint, ₹ lakh (LPA)", fmt="{:.1f}",
     highlight=["Data Analyst", "Data Scientist", "Data Engineer"])
fig.text(0.01, 0.01, "Only postings that disclose salary; roles with ≥100 salary-disclosed postings", fontsize=8, color=GREY)
fig.tight_layout(rect=(0, 0.03, 1, 1)); fig.savefig("outputs/charts/04_salary_by_role.png", dpi=160); plt.close(fig)

# C5 data salary by experience
q14 = R["Q14"]
order = ["Fresher (0-1)", "Junior (1-3)", "Mid (3-6)", "Senior (6-10)", "Lead / Expert (10+)"]
fig, ax = plt.subplots(figsize=(8, 4.8))
for role, col in [("Data Analyst", ORANGE), ("Data Engineer", NAVY), ("Data Scientist", TEAL)]:
    d = q14[q14["role_family"] == role].set_index("experience_level").reindex(order)
    ax.plot(order, d["avg_salary_lpa"], marker="o", lw=2.2, color=col, label=role)
ax.set_title("Average salary (LPA) by experience level: data roles", loc="left", color=NAVY)
ax.set_ylabel("LPA"); ax.legend(frameon=False)
plt.setp(ax.get_xticklabels(), rotation=15)
fig.text(0.01, 0.01, "Cells with <5 salary-disclosed postings omitted", fontsize=8, color=GREY)
fig.tight_layout(rect=(0, 0.03, 1, 1)); fig.savefig("outputs/charts/05_data_salary_by_experience.png", dpi=160); plt.close(fig)

# C6 skill premium
q8 = R["Q08"].head(12).iloc[::-1]
fig, ax = plt.subplots(figsize=(8, 5))
ax.barh(q8["skill"], q8["premium_pct"], color=ORANGE, height=0.68)
for i, (v, n) in enumerate(zip(q8["premium_pct"], q8["n_with_salary"])):
    ax.text(v + 1, i, f"+{v:.0f}%  (n={n})", va="center", fontsize=8)
ax.set_xlim(0, q8["premium_pct"].max() * 1.25)
ax.set_title("Skill salary premium vs avg data-role salary", loc="left", color=NAVY)
ax.set_xlabel("% above average salary of data roles")
fig.tight_layout(); fig.savefig("outputs/charts/06_skill_salary_premium.png", dpi=160); plt.close(fig)

# C7 work mode
q11 = R["Q11"].pivot(index="role_group", columns="work_mode", values="pct_within_group").fillna(0)
q11 = q11.reindex(["Data roles", "Other tech roles", "Non-tech roles"])[["On-site", "Hybrid", "Remote"]]
fig, ax = plt.subplots(figsize=(8, 3.6))
left = np.zeros(len(q11))
for col, c in zip(q11.columns, [NAVY, TEAL, ORANGE]):
    ax.barh(q11.index, q11[col], left=left, color=c, label=col, height=0.55)
    for i, (l, v) in enumerate(zip(left, q11[col])):
        if v > 4: ax.text(l + v / 2, i, f"{v:.0f}%", ha="center", va="center", color="white", fontsize=9, fontweight="bold")
    left += q11[col].values
ax.invert_yaxis(); ax.set_xlim(0, 100)
ax.set_title("Work mode mix", loc="left", color=NAVY); ax.legend(ncol=3, frameon=False, loc="lower center", bbox_to_anchor=(0.5, -0.32))
fig.tight_layout(); fig.savefig("outputs/charts/07_work_mode.png", dpi=160); plt.close(fig)

# ------------------------------------------------------------------ DASHBOARD PREVIEW (mimics the Power BI page)
fig = plt.figure(figsize=(16, 9.5), facecolor=LIGHT)
gs = fig.add_gridspec(4, 12, left=0.115, right=0.98, top=0.88, bottom=0.10, hspace=0.8, wspace=1.9)
fig.text(0.02, 0.945, "India Job Market Intelligence 2025", fontsize=22, fontweight="bold", color=NAVY)
fig.text(0.02, 0.912, f"{len(jobs):,} cleaned job postings  •  Python + SQL + Power BI  •  Snapshot: 3 Oct 2025", fontsize=10.5, color=GREY)

kpis = [
    ("Job postings", f"{len(jobs):,}"),
    ("Data-role postings", f"{n_data:,}"),
    ("Companies hiring", f"{jobs['company_clean'].nunique():,}"),
    ("Salary disclosed", f"{jobs['salary_disclosed'].mean()*100:.0f}%"),
    ("Median data-role salary", f"₹{data_jobs['salary_mid_lpa'].median():.1f} LPA"),
    ("Top hiring city", jobs[jobs.city_group != 'Remote / Unspecified']['city_group'].value_counts().index[0]),
]
for i, (lab, val) in enumerate(kpis):
    ax = fig.add_subplot(gs[0, i * 2:(i + 1) * 2])
    ax.set_facecolor("white"); ax.set_xticks([]); ax.set_yticks([])
    for s in ax.spines.values(): s.set_visible(False)
    ax.text(0.5, 0.58, val, ha="center", va="center", fontsize=19, fontweight="bold", color=NAVY if i != 1 else ORANGE)
    ax.text(0.5, 0.18, lab, ha="center", va="center", fontsize=9.5, color=GREY)

def card(ax):
    ax.set_facecolor("white")

ax1 = fig.add_subplot(gs[1:3, 0:4]); card(ax1)
q5s = R["Q05"].head(10)
hbar(ax1, q5s["skill"], q5s["pct_of_data_jobs"], title="Top skills in data roles (% of postings)", fmt="{:.0f}%",
     highlight=["Python", "SQL", "Power BI", "Excel", "Tableau"])
ax1.set_xlabel("")

ax2 = fig.add_subplot(gs[1:3, 4:8]); card(ax2)
q13s = R["Q13"].query("role_family in ['Data Engineer','Data Scientist','Data Analyst','Software Engineer','Full Stack Developer','DevOps / Cloud','QA / Testing','Product / Project Mgmt']").sort_values("median_lpa", ascending=False)
hbar(ax2, q13s["role_family"], q13s["median_lpa"], color=NAVY, title="Median salary by role (LPA)", fmt="{:.1f}",
     highlight=["Data Analyst", "Data Scientist", "Data Engineer"])
ax2.set_xlabel("")

ax3 = fig.add_subplot(gs[1:3, 8:12]); card(ax3)
q10 = R["Q10"].head(8)
hbar(ax3, q10["city"], q10["data_job_postings"], color=TEAL, title="Data jobs by city", fmt="{:,.0f}")
ax3.set_xlabel("")

ax4 = fig.add_subplot(gs[3, 0:6]); card(ax4)
q17 = R["Q17"]
q17 = q17[q17["experience_level"] != "Not specified"]
ax4.bar(q17["experience_level"], q17["pct"], color=ORANGE, width=0.6)
for i, v in enumerate(q17["pct"]): ax4.text(i, v + 0.8, f"{v:.0f}%", ha="center", fontsize=9)
ax4.set_title("Data-role demand by experience", loc="left", color=NAVY); ax4.set_ylim(0, q17["pct"].max() * 1.2)
plt.setp(ax4.get_xticklabels(), fontsize=8)

ax5 = fig.add_subplot(gs[3, 6:12]); card(ax5)
for role, col in [("Data Analyst", ORANGE), ("Data Engineer", NAVY), ("Data Scientist", TEAL)]:
    d = q14[q14["role_family"] == role].set_index("experience_level").reindex(order)
    ax5.plot(range(len(order)), d["avg_salary_lpa"], marker="o", lw=2.2, color=col, label=role)
ax5.set_xticks(range(len(order))); ax5.set_xticklabels(["0-1", "1-3", "3-6", "6-10", "10+"])
ax5.set_xlabel("Years of experience"); ax5.set_title("Avg salary (LPA) by experience", loc="left", color=NAVY)
ax5.legend(frameon=False, fontsize=8, ncol=3, loc="upper left")
fig.text(0.98, 0.012, "Preview rendered with matplotlib – the interactive version lives in Power BI Desktop (see Build Guide).",
         ha="right", fontsize=8, color=GREY)
fig.savefig("outputs/charts/00_dashboard_preview.png", dpi=150, facecolor=LIGHT); plt.close(fig)

# ------------------------------------------------------------------ EXCEL INSIGHTS WORKBOOK
wb = Workbook()
HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF", size=10)
BODY = Font(name="Arial", size=10)
THIN = Side(style="thin", color="D1D5DB")


def write_table(ws, df, start_row=1, start_col=1, number_formats=None):
    number_formats = number_formats or {}
    for j, col in enumerate(df.columns):
        c = ws.cell(row=start_row, column=start_col + j, value=col)
        c.font = HEADER_FONT; c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for i, row in enumerate(df.itertuples(index=False), start=1):
        for j, v in enumerate(row):
            if isinstance(v, (np.integer,)): v = int(v)
            elif isinstance(v, (np.floating,)): v = None if np.isnan(v) else float(v)
            elif pd.isna(v) if not isinstance(v, (list, dict)) else False: v = None
            c = ws.cell(row=start_row + i, column=start_col + j, value=v)
            c.font = BODY; c.border = Border(bottom=THIN)
            nf = number_formats.get(df.columns[j])
            if nf: c.number_format = nf
    for j, col in enumerate(df.columns):
        width = max(len(str(col)), *(len(str(x)) for x in df[col].head(200))) + 3
        ws.column_dimensions[get_column_letter(start_col + j)].width = min(max(width, 11), 48)
    ws.row_dimensions[start_row].height = 30
    ws.freeze_panes = ws.cell(row=start_row + 1, column=start_col)


# Summary sheet
ws = wb.active; ws.title = "Summary"
ws["A1"] = "India Job Market Intelligence 2025 – Key Findings"; ws["A1"].font = Font(name="Arial", bold=True, size=16, color="1F3A5F")
ws["A2"] = "Source: Kaggle Indian Job Market Dataset 2025 (cleaned with Python; analysed with SQL). Snapshot date: 3-Oct-2025."
ws["A2"].font = Font(name="Arial", italic=True, size=9, color="6B7280")
top5 = R["Q05"].head(5)["skill"].tolist()
da = R["Q13"].set_index("role_family")
facts = [
    ("Raw records", log["raw_rows"]),
    ("Duplicates / reposts removed", log["rows_removed_total"]),
    ("Clean records", log["clean_rows"]),
    ("Distinct raw job titles → standardized", f'{log["distinct_titles_raw"]:,} → {log["distinct_titles_standard"]:,}'),
    ("Salary disclosed (valid)", f'{log["salary_disclosed_pct"]}%'),
    ("Data-role postings (Analyst + Scientist + Engineer)", log["data_roles"]),
    ("Top 5 skills in data roles", ", ".join(top5)),
    ("Median salary: Data Engineer (LPA)", float(da.loc["Data Engineer", "median_lpa"])),
    ("Median salary: Data Scientist (LPA)", float(da.loc["Data Scientist", "median_lpa"])),
    ("Median salary: Data Analyst (LPA)", float(da.loc["Data Analyst", "median_lpa"])),
]
ws["A4"], ws["B4"] = "Metric", "Value"
for c in ("A4", "B4"): ws[c].font = HEADER_FONT; ws[c].fill = HEADER_FILL
for i, (k, v) in enumerate(facts, start=5):
    ws.cell(row=i, column=1, value=k).font = BODY
    c = ws.cell(row=i, column=2, value=v); c.font = BODY
    if isinstance(v, int): c.number_format = "#,##0"
    c.alignment = Alignment(horizontal="left")
ws.column_dimensions["A"].width = 52; ws.column_dimensions["B"].width = 60
ws["A17"] = "Tabs: each sheet below is the output of one SQL query in 02_analysis_queries.sql (Q-number in tab name)."
ws["A17"].font = Font(name="Arial", size=9, color="6B7280")

sheet_names = {
    "Q02": "Q02 Role demand", "Q04": "Q04 Top skills all", "Q05": "Q05 Top skills data", "Q06": "Q06 Skills by data role",
    "Q07": "Q07 SQL co-skills", "Q08": "Q08 Skill premium", "Q09": "Q09 City demand", "Q10": "Q10 Data jobs by city",
    "Q11": "Q11 Work mode", "Q13": "Q13 Salary by role", "Q14": "Q14 Data salary by exp", "Q15": "Q15 Salary bands",
    "Q16": "Q16 Tech salary by city", "Q17": "Q17 Data exp demand", "Q18": "Q18 Top data employers", "Q19": "Q19 Skills by exp",
}
for q, nm in sheet_names.items():
    s = wb.create_sheet(nm)
    df = R[q].copy()
    nf = {c: "0.00" for c in df.columns if any(k in c for k in ["lpa", "pct", "premium", "rating"])}
    nf.update({c: "#,##0" for c in df.columns if c in ("job_postings", "data_job_postings", "tech_jobs", "data_jobs", "all_jobs")})
    write_table(s, df, number_formats=nf)

# native Excel bar charts on two sheets
s = wb["Q05 Top skills data"]
ch = BarChart(); ch.type = "bar"; ch.style = 10; ch.title = "Top skills in data roles (% of postings)"
ch.add_data(Reference(s, min_col=3, min_row=1, max_row=16), titles_from_data=True)
ch.set_categories(Reference(s, min_col=1, min_row=2, max_row=16))
ch.y_axis.title = "% of data postings"; ch.legend = None; ch.height = 10; ch.width = 18
ch.x_axis.scaling.orientation = "maxMin"
s.add_chart(ch, "G2")

s = wb["Q13 Salary by role"]
ch = BarChart(); ch.type = "bar"; ch.style = 10; ch.title = "Median salary by role family (LPA)"
n = len(R["Q13"])
ch.add_data(Reference(s, min_col=4, min_row=1, max_row=min(n + 1, 16)), titles_from_data=True)
ch.set_categories(Reference(s, min_col=1, min_row=2, max_row=min(n + 1, 16)))
ch.legend = None; ch.height = 11; ch.width = 18; ch.x_axis.scaling.orientation = "maxMin"
s.add_chart(ch, "I2")

# cleaned data sample + data dictionary
sample = jobs.sample(2000, random_state=42)[["job_id", "title", "title_standard", "role_family", "company_clean", "city", "city_tier",
        "work_mode", "experience_level", "salary_min_inr", "salary_max_inr", "salary_mid_lpa", "salary_band", "n_skills"]]
write_table(wb.create_sheet("Clean data sample"), sample, number_formats={"salary_min_inr": "#,##0", "salary_max_inr": "#,##0"})

dd = pd.DataFrame([
    ("job_id", "Unique Naukri job id (string)", "raw jobId"),
    ("title / title_standard", "Original title / lower-cased, bracket-stripped, seniority-stripped, Title Case", "regex normalisation"),
    ("role_family", "26 rule-based role families (e.g. Data Analyst, Data Engineer, Software Engineer)", "title + skills rules"),
    ("is_data_role / is_tech_role", "1 if role_family is Data Analyst/Scientist/Engineer; tech = data + software/dev/devops/QA/security", "derived"),
    ("city / city_group / city_tier", "Primary city after alias merge (Bangalore→Bengaluru, Gurgaon/Noida→Delhi NCR, Thane/Navi Mumbai→Mumbai); top-40 else 'Other Cities'; Tier 1/2/Other", "location parsing"),
    ("work_mode", "On-site / Hybrid / Remote, parsed from location prefix", "location parsing"),
    ("salary_min_inr / max / mid", "Annual INR; undisclosed or invalid (<1.2L or >1.5Cr midpoint) set to null; USD ×83", "minimumSalary/maximumSalary"),
    ("salary_mid_lpa", "Salary midpoint in lakh per annum (INR 100,000)", "derived"),
    ("salary_band", "<3, 3-6, 6-10, 10-15, 15-25, 25-50, 50+ LPA", "derived"),
    ("experience_level", "Fresher (0-1) … Lead/Expert (10+) based on minimum experience", "minimumExperience"),
    ("posted_date", "Date parsed from jobId prefix (DDMMYY); fallback = snapshot − 'N days ago'", "jobId / jobUploaded"),
    ("job_skills.skill", "Normalised skill (aliases merged, generic tags dropped)", "tagsAndSkills"),
], columns=["Column", "Description", "Source / method"])
write_table(wb.create_sheet("Data dictionary"), dd)
for ws_ in wb.worksheets:
    ws_.sheet_view.showGridLines = False if ws_.title == "Summary" else True
wb.save("outputs/Job_Market_Insights.xlsx")
print("Excel + charts written")
